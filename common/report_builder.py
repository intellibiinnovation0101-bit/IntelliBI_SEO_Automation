"""
Report builder (common/report_builder.py).

Management-level SEO Walk-In workbook: Summary + Lead Source Trend + Technology
Trend + Lead Type Trend. Each analytical section keeps its TABLE and its CHART
together (chart directly below the table). Performance is shown with a consistent
Green / Yellow / Orange / Red row-level colour scheme.

Colour thresholds (consistent throughout the workbook):
  Growth / Decline %:  Green >= +10% · Yellow 0..+10% · Orange -10..0% · Red < -10%
                       (a brand-new value where Previous = 0 counts as Green)
  Share % (Cur):       Green >= 40% · Yellow 20..40% · Orange 10..20% · Red < 10%
"""
from __future__ import annotations
import collections
import datetime as _dt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.utils import get_column_letter

# ── palette ──────────────────────────────────────────────────────────────────
NAVY = "1B355E"; NAVY2 = "24457A"; LIGHT = "EAF0F8"; ALT = "F5F8FC"
GS_FILL = "FFF2CC"; GS_TEXT = "B7791F"
GREY = "6B7B93"; WHITE = "FFFFFF"
# performance fills (pastel, readable with dark text)
F_GREEN = "D5F5E3"; F_YELLOW = "FCF3CF"; F_ORANGE = "FAE5D3"; F_RED = "F5B7B1"

_thin = Side(style="thin", color="D3DCE8")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEF = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _F(sz=10, b=False, color="1F2A44"):
    return Font(name="Arial", size=sz, bold=b, color=color)


def _fill(c):
    return PatternFill("solid", fgColor=c)


# ── threshold logic (centralised) ────────────────────────────────────────────
def growth_fill(cur, prev):
    if prev == 0:
        return F_YELLOW if cur == 0 else F_GREEN
    g = (cur - prev) / prev
    if g >= 0.10:
        return F_GREEN
    if g >= 0.0:
        return F_YELLOW
    if g >= -0.10:
        return F_ORANGE
    return F_RED


def share_fill(cur, total):
    s = (cur / total) if total else 0.0
    if s >= 0.40:
        return F_GREEN
    if s >= 0.20:
        return F_YELLOW
    if s >= 0.10:
        return F_ORANGE
    return F_RED


# ── data helpers ─────────────────────────────────────────────────────────────
def _cnt(records, key):
    return collections.Counter(r[key] for r in records)


def _gs(records, gs_label):
    return sum(1 for r in records if r["lead_source"] == gs_label)


def _cats_union(cur, prev, key):
    c = _cnt(cur, key); p = _cnt(prev, key)
    cats = list(c.keys())
    for k in p:
        if k not in cats:
            cats.append(k)
    cats.sort(key=lambda k: -(c.get(k, 0) + p.get(k, 0)))
    return cats


def _day_series(records, start, n, key=None):
    """Per-day totals (and per-category counts when key given) for n days from start."""
    out = []
    for i in range(n):
        d = start + _dt.timedelta(days=i)
        drecs = [r for r in records if r["date"] == d]
        counts = collections.Counter(r[key] for r in drecs) if key else None
        out.append((d, len(drecs), counts))
    return out


# ── cell / structural helpers ────────────────────────────────────────────────
def _cell(ws, row, col, val, b=False, al=CEN, num=None, color="1F2A44", bg=None):
    c = ws.cell(row, col, val)
    c.font = _F(9.5, b, color)
    c.alignment = al
    c.border = BORDER
    if num:
        c.number_format = num
    if bg:
        c.fill = _fill(bg)
    return c


def _title(ws, row, text, ncol=6):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row, 1, text)
    c.font = _F(11.5, True, NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 22


def _thead(ws, row, cols, start=1):
    for j, h in enumerate(cols):
        c = ws.cell(row, start + j, h)
        c.font = _F(9.5, True, WHITE)
        c.fill = _fill(NAVY)
        c.alignment = CEN
        c.border = BORDER
    ws.row_dimensions[row].height = 24


def _header_block(ws, period, subtitle, gen_stamp, ncol=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, "IntelliBI  •  SEO Walk-In Analysis")
    c.font = _F(16, True, WHITE); c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    c = ws.cell(2, 1, f"{period.label} Report — {subtitle}")
    c.font = _F(10.5, True, WHITE); c.fill = _fill(NAVY2)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncol)
    c = ws.cell(3, 1,
                f"Current: {period.cur_start.strftime('%d-%b-%Y')} → {period.cur_end.strftime('%d-%b-%Y')}"
                f"   |   Previous (same {period.elapsed_days} day(s)): "
                f"{period.prev_start.strftime('%d-%b-%Y')} → {period.prev_end.strftime('%d-%b-%Y')}"
                f"   |   Generated: {gen_stamp}")
    c.font = _F(8.5, False, GREY)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[3].height = 15


def _kpi(ws, col, row, label, value, sub, accent):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + 1)
    a = ws.cell(row, col, label); a.font = _F(9, True, WHITE); a.fill = _fill(accent); a.alignment = CEN
    b = ws.cell(row + 1, col, value); b.font = _F(20, True, accent); b.alignment = CEN; b.fill = _fill(LIGHT)
    d = ws.cell(row + 2, col, sub); d.font = _F(8, False, GREY); d.alignment = CEN; d.fill = _fill(LIGHT)
    for rr in (row, row + 1, row + 2):
        for cc in (col, col + 1):
            ws.cell(rr, cc).border = BORDER
    ws.row_dimensions[row + 1].height = 28


def _row_fill(ws, row, ncols, bg):
    for c in range(1, ncols + 1):
        ws.cell(row, c).fill = _fill(bg)


# ── comparison table (Current/Previous/Difference[/Share]/Growth, row-coloured) ─
def _comp_table(ws, r, header, rows, color_by, include_share, first_label_col="Category"):
    """rows: list of (label, cur, prev). color_by: 'growth' | 'share'.
    Returns (data_start, data_end, next_row_after_table)."""
    _thead(ws, r, header)
    ncols = len(header)
    r += 1
    start = r
    total_cur = sum(cv for _, cv, _ in rows) or 0
    for label, cur, prev in rows:
        if color_by == "share":
            bg = share_fill(cur, total_cur)
        else:
            bg = growth_fill(cur, prev)
        _cell(ws, r, 1, label, al=LEF, b=False)
        _cell(ws, r, 2, cur)
        _cell(ws, r, 3, prev)
        _cell(ws, r, 4, f"=B{r}-C{r}", num="+0;-0;0")
        if include_share:
            _cell(ws, r, 5, f'=IFERROR(B{r}/SUM(B{start}:B{start + len(rows) - 1}),"")', num="0.0%")
            _cell(ws, r, 6, f'=IFERROR((B{r}-C{r})/C{r},"")', num="0.0%")
        else:
            _cell(ws, r, 5, f'=IFERROR((B{r}-C{r})/C{r},"")', num="0.0%")
        _row_fill(ws, r, ncols, bg)
        r += 1
    return start, r - 1, r


# ── charts ───────────────────────────────────────────────────────────────────
def _clustered(ws, title, cat_ref, data_ref, anchor, span=15, w=15, h=7.2):
    ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.gapWidth = 60
    ch.title = title; ch.style = 10; ch.height = h; ch.width = w
    ch.add_data(data_ref, titles_from_data=True); ch.set_categories(cat_ref)
    ch.y_axis.scaling.min = 0; ch.y_axis.title = "Number of Walk-Ins"
    ws.add_chart(ch, anchor)
    return span


def _stacked(ws, title, cat_ref, data_ref, anchor, span=16, w=18, h=7.6):
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    ch.gapWidth = 40; ch.title = title; ch.style = 10; ch.height = h; ch.width = w
    ch.add_data(data_ref, titles_from_data=True); ch.set_categories(cat_ref)
    ch.y_axis.scaling.min = 0; ch.y_axis.title = "Number of Walk-Ins"
    ch.x_axis.title = "Day"
    ws.add_chart(ch, anchor)
    return span


def _stacked_by_day(ws, title, cat_ref, data_ref, cats, gs_label, anchor,
                    ymax=None, w=8.6, h=8.4):
    """A native single-stack column chart (one stacked bar per day, stacked by
    category). Native stacked columns render identically and cleanly in BOTH
    Excel and Google Sheets — unlike a clustered-+-stacked chart, which Google
    Sheets cannot draw. Two of these (Current, Previous) placed side by side give
    the day-wise Current-vs-Previous distribution comparison. cats -> consistent
    per-category colours; ymax makes the two charts share a Y scale."""
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    ch.gapWidth = 60; ch.title = title; ch.style = 10; ch.height = h; ch.width = w
    ch.add_data(data_ref, titles_from_data=True); ch.set_categories(cat_ref)
    ch.y_axis.scaling.min = 0
    if ymax:
        ch.y_axis.scaling.max = ymax
    ch.y_axis.title = "Walk-Ins"; ch.x_axis.title = "Day"
    for idx, ser in enumerate(ch.series):
        cat = cats[idx] if idx < len(cats) else ""
        ser.graphicalProperties.solidFill = _series_color(cat, idx, gs_label)
    ws.add_chart(ch, anchor)
    return ch


def _lines(ws, title, cat_ref, data_refs, anchor, span=15, w=17, h=7.2):
    lc = LineChart(); lc.title = title; lc.style = 2; lc.height = h; lc.width = w
    for dref in data_refs:
        lc.add_data(dref, titles_from_data=True)
    lc.set_categories(cat_ref)
    for s in lc.series:
        s.marker = Marker(symbol="circle", size=6)
        s.smooth = False
    lc.y_axis.scaling.min = 0; lc.y_axis.title = "Number of Walk-Ins"; lc.x_axis.title = "Day"
    ws.add_chart(lc, anchor)
    return span


# ── clustered + stacked chart (Current vs Previous per date) ─────────────────
# Consistent category colours (same category -> same colour in every bar).
_CAT_PALETTE = ["1B355E", "2E7D32", "C0392B", "6C8EBF", "9B59B6",
                "16A085", "E67E22", "7F8C8D", "34495E", "1ABC9C", "8E44AD"]


def _series_color(cat, idx, gs_label):
    if cat == gs_label:
        return GS_TEXT                      # Google Search keeps its gold identity
    return _CAT_PALETTE[idx % len(_CAT_PALETTE)]


def _clustered_stacked(ws, title, cat_ref, data_ref, cats, gs_label, anchor,
                       span=18, w=22, h=9):
    """Stacked column chart with a 2-LEVEL category axis (outer = date, inner =
    Current/Previous). Excel/LibreOffice render this as two stacked bars per date
    — the clustered-+-stacked comparison. cat_ref must span the two category
    columns (date, period); data_ref the category-count columns (one series each)."""
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    # gapWidth 0 -> the two bars in a date group (Current, Previous) touch; a blank
    # spacer row in the source data creates the gap BETWEEN dates.
    ch.gapWidth = 0; ch.title = title; ch.style = 10; ch.height = h; ch.width = w
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cat_ref)
    ch.y_axis.scaling.min = 0
    ch.y_axis.title = "Number of Walk-Ins"
    # consistent per-category colours (same category -> same colour in every bar)
    for idx, ser in enumerate(ch.series):
        cat = cats[idx] if idx < len(cats) else ""
        ser.graphicalProperties.solidFill = _series_color(cat, idx, gs_label)
    ws.add_chart(ch, anchor)
    return span


# ── SUMMARY ──────────────────────────────────────────────────────────────────
def _summary(wb, period, cur, prev, gs, gen_stamp):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    _header_block(ws, period, "Walk-In Performance & Google Search (SEO)", gen_stamp, ncol=8)
    tot_c, tot_p = len(cur), len(prev)
    gs_c, gs_p = _gs(cur, gs), _gs(prev, gs)

    # Executive Snapshot — Total Walk-Ins only, row coloured by growth.
    # Section heading uses the navy header-band style (white bold on navy),
    # matching the (now removed) KPI header cards.
    r = 4
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    c = ws.cell(r, 1, "Executive Snapshot")
    c.font = _F(9, True, WHITE)          # same font type/size as the removed KPI header
    c.fill = _fill(NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for cc in range(1, 6):
        ws.cell(r, cc).border = BORDER
    ws.row_dimensions[r].height = 24
    r += 1
    ds, de, r = _comp_table(
        ws, r, ["Metric", "Current", "Previous", "Difference", "Growth / Decline %"],
        [("Total Walk-Ins", tot_c, tot_p)], color_by="growth", include_share=False)
    r += 1

    def agg_block(r, heading, key, chart_kind):
        _title(ws, r, heading, ncol=6); r += 1
        rows = [(cat, _cnt(cur, key).get(cat, 0), _cnt(prev, key).get(cat, 0))
                for cat in _cats_union(cur, prev, key)]
        ds, de, r = _comp_table(
            ws, r, [heading.split(" — ")[0].replace(" Trend", ""), "Current", "Previous",
                    "Difference", "Share % (Cur)", "Growth / Decline %"],
            rows, color_by="share", include_share=True)
        # highlight the Google Search label cell on the Lead Source block
        if key == "lead_source":
            for rr in range(ds, de + 1):
                if ws.cell(rr, 1).value == gs:
                    ws.cell(rr, 1).font = _F(9.5, True, GS_TEXT)
        chart_anchor = f"A{r + 1}"
        if chart_kind == "total_vs_gs":
            # helper data placed FAR RIGHT (cols AF:AH) and excluded from the print
            # area, so it never shows in the sheet while still feeding the chart.
            hc = 32
            hr = ds
            ws.cell(hr - 1, hc, "Measure"); ws.cell(hr - 1, hc + 1, "Current"); ws.cell(hr - 1, hc + 2, "Previous")
            ws.cell(hr, hc, "Total Walk-Ins"); ws.cell(hr, hc + 1, tot_c); ws.cell(hr, hc + 2, tot_p)
            ws.cell(hr + 1, hc, "Google Search"); ws.cell(hr + 1, hc + 1, gs_c); ws.cell(hr + 1, hc + 2, gs_p)
            cat_ref = Reference(ws, min_col=hc, max_col=hc, min_row=hr, max_row=hr + 1)
            data_ref = Reference(ws, min_col=hc + 1, max_col=hc + 2, min_row=hr - 1, max_row=hr + 1)
            span = _clustered(ws, "Total vs Google Search — Current vs Previous",
                              cat_ref, data_ref, chart_anchor, w=13, h=6.8)
        else:
            cat_ref = Reference(ws, min_col=1, max_col=1, min_row=ds, max_row=de)
            data_ref = Reference(ws, min_col=2, max_col=3, min_row=ds - 1, max_row=de)
            span = _clustered(ws, heading.replace(" Trend", "") + " — Current vs Previous",
                              cat_ref, data_ref, chart_anchor, w=15, h=7.2)
        return r + 1 + span

    r = agg_block(r, "Lead Source Trend", "lead_source", "total_vs_gs")
    r = agg_block(r, "Technology Trend", "technology", "cluster")
    r = agg_block(r, "Lead Type Trend", "lead_type", "cluster")

    for col, w in {"A": 26, "B": 12, "C": 12, "D": 12, "E": 16, "F": 18}.items():
        ws.column_dimensions[col].width = w
    # Keep the printed/exported area to the visible report; the chart-data helper
    # block sits far right (col AF+) and is left out.
    ws.print_area = f"A1:{get_column_letter(22)}{r + 1}"
    return ws


# ── a trend tab: day-wise total (+line) then day-wise distribution (2 stacked) ─
def _trend_tab(wb, tab_name, period, cur, prev, key, subtitle, gs):
    ws = wb.create_sheet(tab_name)
    ws.sheet_view.showGridLines = False
    _header_block(ws, period, subtitle, gen_stamp=_GEN[0], ncol=10)

    # A. Day-Wise Walk-In Trend (Current vs Previous), row coloured by growth
    r = 4
    _title(ws, r, "Day-Wise Walk-In Trend (Current vs Previous)", ncol=5); r += 1
    _thead(ws, r, ["Date / Day", "Current", "Previous", "Difference", "Growth / Decline %"]); r += 1
    ds = r
    cur_days = _day_series(cur, period.cur_start, period.elapsed_days)
    prev_days = _day_series(prev, period.prev_start, period.elapsed_days)
    for i in range(period.elapsed_days):
        cd, c, _ = cur_days[i]
        pd, p, _ = prev_days[i]
        bg = growth_fill(c, p)
        _cell(ws, r, 1, f"{cd.strftime('%a %d-%b')}", al=LEF)
        _cell(ws, r, 2, c); _cell(ws, r, 3, p)
        _cell(ws, r, 4, f"=B{r}-C{r}", num="+0;-0;0")
        _cell(ws, r, 5, f'=IFERROR((B{r}-C{r})/C{r},"")', num="0.0%")
        _row_fill(ws, r, 5, bg)
        r += 1
    de = r - 1
    _cell(ws, r, 1, "Total", b=True, al=LEF, bg=LIGHT)
    _cell(ws, r, 2, f"=SUM(B{ds}:B{de})", b=True, bg=LIGHT)
    _cell(ws, r, 3, f"=SUM(C{ds}:C{de})", b=True, bg=LIGHT)
    _cell(ws, r, 4, f"=B{r}-C{r}", b=True, bg=LIGHT, num="+0;-0;0")
    _cell(ws, r, 5, f'=IFERROR((B{r}-C{r})/C{r},"")', b=True, bg=LIGHT, num="0.0%")
    r += 2
    cat_ref = Reference(ws, min_col=1, max_col=1, min_row=ds, max_row=de)
    cur_ref = Reference(ws, min_col=2, max_col=2, min_row=ds - 1, max_row=de)
    prev_ref = Reference(ws, min_col=3, max_col=3, min_row=ds - 1, max_row=de)
    span = _lines(ws, "Walk-Ins — Current vs Previous (day-wise)", cat_ref, [cur_ref, prev_ref], f"A{r}")
    r += span

    # B. Day-Wise Distribution — ONE combined Current vs Previous table + ONE
    #    clustered+stacked chart (two stacked bars per date: Current & Previous).
    label = subtitle.replace(" Trend Analysis", "")
    cats = _cats_union(cur, prev, key)
    ncat = len(cats)
    ncols = 4 + ncat        # Date | Period | Comparable Date | Total | cats...

    _title(ws, r, f"{label} Distribution — Day-Wise Current vs Previous", ncol=ncols); r += 1
    _thead(ws, r, ["Date", "Period", "Comparable Date", "Total"] + cats); r += 1
    s = r
    cur_ser = _day_series(cur, period.cur_start, period.elapsed_days, key=key)
    prev_ser = _day_series(prev, period.prev_start, period.elapsed_days, key=key)
    for i in range(period.elapsed_days):
        cd, ctot, ccnt = cur_ser[i]
        pd, ptot, pcnt = prev_ser[i]
        pos = cd.strftime("%a %d-%b")                    # outer axis = CURRENT date
        for period_name, actual, tot, cnt, base_bg in (
                ("Current", cd, ctot, ccnt, WHITE),
                ("Previous", pd, ptot, pcnt, ALT)):
            _cell(ws, r, 1, pos, al=LEF, b=(period_name == "Current"))
            _cell(ws, r, 2, period_name, b=True,
                  color=("1F2A44" if period_name == "Current" else GREY))
            _cell(ws, r, 3, actual.strftime("%a %d-%b"), al=LEF, color=GREY)
            _cell(ws, r, 4, tot, b=True, bg=LIGHT)
            for j, cat in enumerate(cats):
                v = cnt.get(cat, 0) if cnt else 0
                bg = GS_FILL if (key == "lead_source" and cat == gs) else base_bg
                _cell(ws, r, 5 + j, v, bg=bg)
            r += 1
    e = r - 1
    r += 1

    # Chart data helper (off to the right, excluded from print): Current & Previous
    # bars per date. A BLANK SPACER row between date groups + gapWidth 0 makes the
    # two bars of a date TOUCH while dates stay apart. Each bar carries ITS OWN date
    # in the axis column — the Current bar the current date, the Previous bar the
    # PREVIOUS comparable date — so the x-axis no longer repeats the current date
    # under both bars. Period (Curr/Prev) stays as the second category level.
    hcol = 32                                 # far right, clear of the chart & print area
    hcat0 = hcol + 2                          # first category-count column
    hhdr = s - 1
    ws.cell(hhdr, hcol, "Date"); ws.cell(hhdr, hcol + 1, "Period")
    for j, cat in enumerate(cats):
        ws.cell(hhdr, hcat0 + j, cat)
    hr = s
    for i in range(period.elapsed_days):
        cd, ctot, ccnt = cur_ser[i]
        pd, ptot, pcnt = prev_ser[i]
        # compact X-axis labels ONLY (chart), e.g. "M|31Aug"; the visible table
        # keeps the full "Mon 31-Aug" / "Current" / "Previous" labels. The Current
        # bar shows the current date, the Previous bar shows its comparable date.
        cd_short = cd.strftime("%a")[0] + "|" + cd.strftime("%d%b")
        pd_short = pd.strftime("%a")[0] + "|" + pd.strftime("%d%b")
        for date_short, period_short, cnt in (
                (cd_short, "Curr", ccnt), (pd_short, "Prev", pcnt)):
            ws.cell(hr, hcol, date_short)
            ws.cell(hr, hcol + 1, period_short)
            for j, cat in enumerate(cats):
                ws.cell(hr, hcat0 + j, (cnt.get(cat, 0) if cnt else 0))
            hr += 1
        if i < period.elapsed_days - 1:       # blank spacer between date groups
            hr += 1
    hend = hr - 1

    cat_ref = Reference(ws, min_col=hcol, max_col=hcol + 1, min_row=s, max_row=hend)
    data_ref = Reference(ws, min_col=hcat0, max_col=hcat0 + ncat - 1, min_row=hhdr, max_row=hend)
    span = _clustered_stacked(
        ws, f"{label} Distribution — Current vs Previous (per date)",
        cat_ref, data_ref, cats, gs, f"A{r}")
    r += span

    widths = {"A": 14, "B": 11, "C": 14, "D": 8}
    for j in range(ncat):
        widths[get_column_letter(5 + j)] = 13
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    # Print/PDF/email area covers the tables + charts (through the chart's right
    # edge) but stops well before the off-to-the-right helper block at col 32.
    last_row = max(r, hend) + 1
    ws.print_area = f"A1:{get_column_letter(22)}{last_row}"
    return ws


_GEN = [""]   # module-local generation stamp for _trend_tab header


def build_workbook(period, cur, prev, gs_label, gen_stamp, sample=False):
    _GEN[0] = gen_stamp
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    _summary(wb, period, cur, prev, gs_label, gen_stamp)
    _trend_tab(wb, "Lead Source Trend", period, cur, prev, "lead_source",
               "Lead Source Trend Analysis", gs_label)
    _trend_tab(wb, "Technology Trend", period, cur, prev, "technology",
               "Technology Trend Analysis", gs_label)
    _trend_tab(wb, "Lead Type Trend", period, cur, prev, "lead_type",
               "Lead Type Trend Analysis", gs_label)
    for ws_ in wb.worksheets:
        ws_.page_setup.orientation = "landscape"
        ws_.page_setup.fitToWidth = 1
        ws_.page_setup.fitToHeight = 0
        ws_.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    return wb
