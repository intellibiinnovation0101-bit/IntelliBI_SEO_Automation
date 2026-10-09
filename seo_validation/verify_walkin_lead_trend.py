"""
Verification of the Weekly / Monthly Walk-In Lead Performance vs target
(common/walkin_targets.py, the e-mail section in common/email_utils.py and the
"Weekly / Monthly Lead Trend" tab in common/report_builder.py).

Run from the project root (no Google access, synthetic walk-ins only):
    python seo_validation\\verify_walkin_lead_trend.py
"""
import datetime as dt
import io
import os
import random
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "common"))
import paths                                                    # noqa: E402
paths.LOGS_DIR = paths.Path(tempfile.mkdtemp(prefix="seo_logs_"))   # never write the real logs/
import openpyxl                                                 # noqa: E402
import report_periods as RP                                     # noqa: E402
import walkin_targets as WT                                     # noqa: E402
import report_builder as RB                                     # noqa: E402
import email_utils as EU                                        # noqa: E402

FAIL = []


def check(label, got, want):
    ok = got == want
    print(f"[{'pass' if ok else 'FAIL'}] {label}: {got!r}" + ("" if ok else f"  (expected {want!r})"))
    if not ok:
        FAIL.append(label)


D = dt.date
# ── 1. targets ───────────────────────────────────────────────────────────────
print("== 1. Targets ==")
check("31-day month week = 7 × 70/31", round(WT.target_for_days(D(2026, 10, 5), D(2026, 10, 11), 70), 3),
      round(7 * 70 / 31, 3))
check("30-day month week = 7 × 70/30", round(WT.target_for_days(D(2026, 9, 7), D(2026, 9, 13), 70), 3),
      round(7 * 70 / 30, 3))
check("February week = 7 × 70/28 = 17.5", round(WT.target_for_days(D(2027, 2, 8), D(2027, 2, 14), 70), 3), 17.5)
check("week across a month end takes each day from its own month (Mon 28-Sep … Sun 04-Oct)",
      round(WT.target_for_days(D(2026, 9, 28), D(2026, 10, 4), 70), 4),
      round(3 * 70 / 30 + 4 * 70 / 31, 4))
check("daily targets of a whole month add back to the monthly target (Oct, Feb, Sep)",
      [round(WT.target_for_days(D(2026, m, 1), D(2026, m, n), 70), 6) for m, n in ((10, 31), (2, 28), (9, 30))],
      [70.0, 70.0, 70.0])
check("status: > target green 'Above', = target red 'Below', < target red",
      [WT.status(17, 16.3), WT.status(16.3, 16.3), WT.status(10, 16.3)], [WT.ABOVE, WT.BELOW, WT.BELOW])

# ── 2. synthetic walk-ins ────────────────────────────────────────────────────
rnd = random.Random(7)
recs = []
d = D(2025, 9, 1)
while d <= D(2026, 10, 31):
    for _ in range(rnd.randint(0, 5)):
        recs.append({"date": d, "lead_source": "Google Search", "technology": "Power BI",
                     "lead_type": "Student", "mobile": "", "name": "x", "tab": "Walk-In New"})
    d += dt.timedelta(days=1)
recs.append({"date": None, "lead_source": "", "technology": "", "lead_type": "", "mobile": "",
             "name": "no date", "tab": "Walk-In New"})        # never counted (same as the report)
cnt = lambda a, b: sum(1 for r in recs if r["date"] and a <= r["date"] <= b)

print("\n== 2. Weekly trend (as of Thu 08-Oct-2026) ==")
pw = RP.weekly(D(2026, 10, 8))
tw = WT.trend_for_period(recs, pw, 70, n_email=5, n_tab=11)
check("12 weeks in the tab, 6 in the e-mail, the e-mail = the tab's last 6",
      (len(tw["tab"]), len(tw["email"]), tw["email"] == tw["tab"][-6:]), (12, 6, True))
check("every week is Monday → Sunday, consecutive, oldest first",
      all(x["start"].weekday() == 0 and (x["end"] - x["start"]).days == 6 for x in tw["tab"])
      and all((b["start"] - a["start"]).days == 7 for a, b in zip(tw["tab"], tw["tab"][1:])), True)
check("first / current week", (tw["tab"][0]["start"], tw["tab"][-1]["start"]), (D(2026, 7, 20), D(2026, 10, 5)))
cur = tw["tab"][-1]
check("current week: counted Mon 05-Oct → Thu 08-Oct = the report's Current total",
      (cur["in_progress"], cur["days_elapsed"], cur["actual"]),
      (True, 4, cnt(pw.cur_start, pw.cur_end)))
check("… full-week target (not pro-rata) and target to date for context",
      (cur["target"], cur["target_to_date"]), (round(7 * 70 / 31, 1), round(4 * 70 / 31, 1)))
check("completed weeks counted in full, never twice (sum of 12 weeks = all walk-ins in the range)",
      sum(x["actual"] for x in tw["tab"]), cnt(D(2026, 7, 20), D(2026, 10, 8)))
check("historical reference date (WEEKLY_REFERENCE_DATE = Sun 27-Sep-2026): that week is complete",
      (lambda t: (t["tab"][-1]["start"], t["tab"][-1]["in_progress"], t["tab"][-1]["actual"]))(
          WT.trend_for_period(recs, RP.weekly(D(2026, 9, 27)), 70)),
      (D(2026, 9, 21), False, cnt(D(2026, 9, 21), D(2026, 9, 27))))

print("\n== 3. Monthly trend ==")
pm = RP.monthly(D(2026, 10, 8))
tm = WT.trend_for_period(recs, pm, 70, n_email=5, n_tab=11)
check("12 months (Nov-2025 … Oct-2026), 6 in the e-mail",
      ([WT.period_label(x) for x in (tm["tab"][0], tm["tab"][-1])], len(tm["email"])), (["Nov-2025", "Oct-2026"], 6))
check("every month's target = 70, current month in progress (counted to 08-Oct)",
      ({x["target"] for x in tm["tab"]}, tm["tab"][-1]["in_progress"], tm["tab"][-1]["actual"]),
      ({70.0}, True, cnt(D(2026, 10, 1), D(2026, 10, 8))))
check("a past month chosen (MONTHLY_MONTH = 9) -> Sep-2026 complete",
      (lambda t: (WT.period_label(t["tab"][-1]), t["tab"][-1]["in_progress"]))(
          WT.trend_for_period(recs, RP.monthly(D(2026, 9, 30)), 70)), ("Sep-2026", False))
check("Manual report -> no trend", WT.trend_for_period(recs, RP.manual(D(2026, 9, 1), D(2026, 9, 15))), None)

# ── 4. workbook ──────────────────────────────────────────────────────────────
print("\n== 4. Workbook tab ==")
_win = lambda p: ([r for r in recs if r["date"] and p.cur_start <= r["date"] <= p.cur_end],
                  [r for r in recs if r["date"] and p.prev_start <= r["date"] <= p.prev_end])
for p, t, name in ((pw, tw, "Weekly Lead Trend"), (pm, tm, "Monthly Lead Trend")):
    c, pv = _win(p)
    wb = RB.build_workbook(p, c, pv, "Google Search", "08-Oct-2026 07:00 PM", trend=t)
    check(f"{p.label}: existing tabs unchanged, new tab last",
          wb.sheetnames, ["Summary", "Lead Source Trend", "Technology Trend", "Lead Type Trend", name])
    ws = wb[name]
    hdr = next(r for r in range(1, 40) if ws.cell(r, 1).value == "#")
    rows = [[ws.cell(r, c).value for c in range(1, 7)] for r in range(hdr + 1, hdr + 13)]
    check(f"{p.label}: 12 rows = the trend (actual, target, status text)",
          [(r[2], r[3], r[5].split(" · ")[0]) for r in rows],
          [(x["actual"], x["target"], x["status"]) for x in t["tab"]])
    check(f"{p.label}: row colours green above / red at-or-below",
          [ws.cell(hdr + 1 + i, 3).fill.fgColor.rgb[-6:] for i in range(12)],
          [RB.F_GREEN if x["above"] else RB.F_RED for x in t["tab"]])
    check(f"{p.label}: current row marked", ("▶ current" in rows[-1][1], "In progress" in rows[-1][5]), (True, True))
    tot = [ws.cell(hdr + 13, c).value for c in range(2, 5)]
    check(f"{p.label}: completed total row = sum of the 11 completed rows",
          tot[1], sum(x["actual"] for x in t["tab"][:-1]))
    ch = ws._charts
    check(f"{p.label}: one combo chart — 4 bar series (green/red, current lighter) + dashed target line",
          (len(ch), [s.graphicalProperties.solidFill.srgbClr.val if hasattr(s.graphicalProperties.solidFill.srgbClr, "val")
                     else s.graphicalProperties.solidFill.srgbClr for s in ch[0].series],
           len(ch[0]._charts) if hasattr(ch[0], "_charts") else None),
          (1, [RB.C_ABOVE, RB.C_BELOW, RB.C_ABOVE_CUR, RB.C_BELOW_CUR], 2))
    helper = [[ws.cell(hdr + 1 + i, 32 + j).value for j in range(6)] for i in range(12)]
    check(f"{p.label}: chart data — exactly one bar per period, in the right colour slot, target alongside",
          all(sum(v is not None for v in h[1:5]) == 1 and h[5] == x["target"]
              and h[1 + ((2 if x["above"] else 3) if x["in_progress"] else (0 if x["above"] else 1))] == x["actual"]
              for h, x in zip(helper, t["tab"])), True)
    check(f"{p.label}: current period label flagged '*' on the chart axis", helper[-1][0].endswith("*"), True)
    buf = io.BytesIO()
    wb.save(buf)
    check(f"{p.label}: workbook saves and re-opens", name in openpyxl.load_workbook(io.BytesIO(buf.getvalue())).sheetnames, True)
c, pv = _win(RP.manual(D(2026, 9, 1), D(2026, 9, 15)))
check("Manual workbook: no Lead Trend tab (unchanged)",
      RB.build_workbook(RP.manual(D(2026, 9, 1), D(2026, 9, 15)), c, pv, "Google Search", "x").sheetnames,
      ["Summary", "Lead Source Trend", "Technology Trend", "Lead Type Trend"])

# ── 5. e-mail ────────────────────────────────────────────────────────────────
print("\n== 5. E-mail ==")
c, pv = _win(pw)
html_w = EU.build_body(pw, c, pv, "Google Search", "https://x", "08-Oct-2026", trend=tw)
check("weekly e-mail: section title",
      "Weekly Walk-In Lead Performance &ndash; Last 5 Completed Weeks + Current Week" in html_w, True)
check("weekly e-mail: the 6 weeks' figures = the tab's last 6 rows",
      all(f"{x['actual']} / {x['target']:g}" in html_w for x in tw["tab"][-6:]), True)
check("weekly e-mail: current week flagged 'In progress', newest first",
      (html_w.index("In progress") < html_w.index(WT.period_label(tw["tab"][-2]))), True)
c, pv = _win(pm)
html_m = EU.build_body(pm, c, pv, "Google Search", "https://x", "08-Oct-2026", trend=tm)
check("monthly e-mail: section title + 70 target",
      ("Monthly Walk-In Lead Performance &ndash; Last 5 Completed Months + Current Month" in html_m,
       "Monthly target 70" in html_m), (True, True))
base_dir = os.environ.get("SEO_BASE_DIR")
if base_dir:                        # optional: compare with the pre-change module
    import importlib.util
    spec = importlib.util.spec_from_file_location("eu_old", os.path.join(base_dir, "common", "email_utils.py"))
    old = importlib.util.module_from_spec(spec); spec.loader.exec_module(old)
    check("no trend (Manual / old callers): e-mail body identical to before",
          EU.build_body(pw, c, pv, "Google Search", "https://x", "s"), old.build_body(pw, c, pv, "Google Search", "https://x", "s"))

print("\nALL CHECKS PASSED" if not FAIL else f"\n{len(FAIL)} CHECK(S) FAILED: {FAIL}")
sys.exit(1 if FAIL else 0)
