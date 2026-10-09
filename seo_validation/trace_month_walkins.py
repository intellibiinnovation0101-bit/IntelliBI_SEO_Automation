"""
Trace ONE month's Walk-In count back to the source rows, with the evidence of
how every Timestamp was read (read-only investigation).

Uses the report's OWN data path — common/walkin_data.records_from_rows (same
tabs, same column mapping, same DATE RULE, same normalisation) and
walkin_data.dedupe — and the Monthly Lead Trend's month rule (date inside the
calendar month; undated rows never counted).

READ-ONLY: reads the source Google Sheet (spreadsheets.readonly) and writes one
investigation workbook to output/. Changes no code, report, sheet or setting.

    python seo_validation\\trace_month_walkins.py                 # May-2026
    python seo_validation\\trace_month_walkins.py --month 2026-06

TRACE EVIDENCE FORMAT  (output\\_trace - Walk-Ins <Mon-YYYY>.xlsx)
  Summary            Spreadsheet locale + time zone. Per tab: the Timestamp
                     column, how many cells are real DATE VALUES vs TEXT vs blank,
                     the date formats in use, the most common display shapes,
                     and the month's count under each reading:
                       report (DATE RULE)        <- what every report now uses
                       previous report logic     (display text, before the fix)
                       date-value cells only     (what a Sheets date filter shows)
                       read as DD/MM from display (what someone reading the
                                                  displayed text day-first counts)
                       each other date-like column in the tab
  Counted <Mon>      One row per counted lead: Lead Name, Mobile Number, Visit
                     Date, Sheet shows (displayed text), Cell holds (date value /
                     text), Date rule (the DATE RULE step), Lead Source (sheet and
                     normalised), Technology, Lead Type, Source Sheet, Sheet Row,
                     Flags (repeat mobile, no mobile, out of sequence, …).
  Month by month     Every month in the data: corrected vs previous count per tab
                     and in total, with the difference — historical validation.
  Text dates         Every Timestamp typed as TEXT (any month): tab, row, raw
                     text, rule applied, resulting date.
  Other date columns Rows of the month whose other date-like column falls in a
                     different month than the Timestamp.
"""
from __future__ import annotations
import argparse
import calendar
import collections
import datetime as dt
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "common"))

import paths                       # noqa: E402
import config_loader as cfg        # noqa: E402
import walkin_data as WD           # noqa: E402
import walkin_targets as WT        # noqa: E402

# other columns that hold a DATE (e.g. "Date of Visit", "Next Follow-Up Date") —
# not time slots / durations such as "Preferred Time Slot"
_DATE_HDR = re.compile(r"date|visit", re.I)
_DISP_DMY = re.compile(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})")


def _shape(v):
    return "(blank)" if v in (None, "") else re.sub(r"\d", "9", str(v))[:24]


def _month_of(d):
    return dt.date(d.year, d.month, 1) if d else None


def _other_date(v):
    """Date of an 'other' date-like cell, read by the same DATE RULE (no context)."""
    d, _r, _s = WD._resolve_dates([v])
    return d[0]


def trace(tabs: dict, month_start: dt.date):
    """tabs = {tab: {"val": rows read by value (SERIAL_NUMBER), "disp": rows as
    displayed (FORMATTED_STRING), "formats": Counter of number-format patterns or
    None}}. Returns a result dict used by write_xlsx / main."""
    month_end = month_start.replace(day=calendar.monthrange(month_start.year, month_start.month)[1])
    in_month = lambda d: d is not None and month_start <= d <= month_end
    fmap_all = WD._mapping()
    new_all, old_all, info, per_tab = [], [], {}, {}
    for tab, src in tabs.items():
        which = "new" if "new" in tab.lower() else "old"
        diag = []
        recs = WD.records_from_rows(tab, src["val"], which, diag)       # the report's own function
        header = src["val"][0] if src["val"] else []
        di = WD._col_index(header, fmap_all[which]["date"])
        si = WD._col_index(header, fmap_all[which]["lead_source"])
        disp_rows = src["disp"]
        others = [(i, h) for i, h in enumerate(header) if i != di and _DATE_HDR.search(str(h or ""))]
        types = collections.Counter()
        tab_old = []
        for rec, dg in zip(recs, diag):
            drow = disp_rows[dg["row"] - 1] if dg["row"] - 1 < len(disp_rows) else []
            shown = drow[di] if (di is not None and di < len(drow)) else ""
            old_d = WD.parse_date(shown)                                 # previous logic: display text
            kind = ("blank" if dg["rule"] == "blank" else "date value" if dg["rule"] == "date value"
                    else "unparsed" if dg["rule"] == "unparsed" else "text")
            types[kind] += 1
            vrow = src["val"][dg["row"] - 1]
            oth = {h: _other_date(vrow[i] if i < len(vrow) else "") for i, h in others}
            info[id(rec)] = {**dg, "shown": shown, "kind": kind, "old": old_d, "other": oth,
                             "src": (vrow[si] if si is not None and si < len(vrow) else "")}
            old_rec = dict(rec); old_rec["date"] = old_d
            info[id(old_rec)] = info[id(rec)]
            tab_old.append(old_rec)
        new_all += recs
        old_all += tab_old
        per_tab[tab] = {"types": types, "formats": src.get("formats"),
                        "shapes": collections.Counter(_shape(info[id(r)]["shown"]) for r in recs).most_common(6),
                        "others": [h for _, h in others], "recs": recs}

    kept, removed = WD.dedupe(new_all)                                   # the report's own de-dup
    kept_old, _ = WD.dedupe(old_all)
    counted = [r for r in kept if in_month(r["date"])]
    assert len(counted) == WT.count_between(kept, month_start, month_end)

    # readings of the month, per tab
    readings = {}
    for tab, pt in per_tab.items():
        ids = {id(r) for r in pt["recs"]}
        kt = [r for r in kept if id(r) in ids]
        rd = collections.OrderedDict()
        rd["report (DATE RULE)"] = sum(1 for r in kt if in_month(r["date"]))
        rd["previous report logic (display text)"] = sum(
            1 for r in kept_old if r["tab"] == tab and in_month(r["date"]))
        rd["date-value cells only (Sheets date filter)"] = sum(
            1 for r in kt if in_month(r["date"]) and info[id(r)]["kind"] == "date value")
        def _dmy(r):
            m = _DISP_DMY.match(str(info[id(r)]["shown"] or ""))
            return bool(m) and int(m.group(2)) == month_start.month and int(m.group(3)) == month_start.year
        rd["read as DD/MM from the displayed text"] = sum(1 for r in kt if _dmy(r))
        for h in pt["others"]:
            rd[f"by other column '{h}'"] = sum(1 for r in kt if in_month(info[id(r)]["other"].get(h)))
        readings[tab] = rd

    # neighbours (sheet order) for out-of-sequence flags
    neighbour = {}
    for pt in per_tab.values():
        lst = pt["recs"]
        for i, r in enumerate(lst):
            prev = next((x["date"] for x in reversed(lst[:i]) if x["date"]), None)
            nxt = next((x["date"] for x in lst[i + 1:] if x["date"]), None)
            neighbour[id(r)] = (prev, nxt)
    mob_month = collections.Counter(r["mobile"] for r in counted if r["mobile"])
    mob_all = collections.defaultdict(list)
    for r in kept:
        if r["mobile"]:
            mob_all[r["mobile"]].append(r)

    rows_out = []
    for k, r in enumerate(sorted(counted, key=lambda x: (x["date"], x["tab"], info[id(x)]["row"])), 1):
        m = info[id(r)]
        prev, nxt = neighbour[id(r)]
        f = []
        if m["kind"] == "text":
            f.append("Timestamp typed as TEXT")
        if "AMBIGUOUS" in m["rule"]:
            f.append(m["rule"])
        if m["old"] != r["date"]:
            f.append(f"previous logic read {m['old']:%d-%b-%Y}" if m["old"] else "previous logic: no date")
        if prev and nxt and not (min(prev, nxt) - dt.timedelta(days=7) <= r["date"] <= max(prev, nxt) + dt.timedelta(days=7)):
            f.append(f"OUT OF SEQUENCE: rows above/below dated {prev:%d-%b-%Y} / {nxt:%d-%b-%Y}")
        for h, od in m["other"].items():
            if od and _month_of(od) != _month_of(r["date"]):
                f.append(f"'{h}' = {od:%d-%b-%Y}")
        if not r["mobile"]:
            f.append("no mobile — never de-duplicated")
        if r["mobile"] and mob_month[r["mobile"]] > 1:
            f.append(f"same mobile {mob_month[r['mobile']]}x this month (different days — each counted)")
        others = [x for x in mob_all.get(r["mobile"], []) if x is not r]
        if others:
            f.append("also a walk-in on " + ", ".join(
                (f"{x['date']:%d-%b-%Y}" if x["date"] else "no date") + f" ({x['tab']})" for x in others[:4]))
        rows_out.append({"#": k, "Lead Name": r["name"], "Mobile Number": r["mobile"], "Visit Date": r["date"],
                         "Sheet shows": m["shown"], "Cell holds": m["kind"], "Date rule": m["rule"],
                         "Lead Source (sheet)": m["src"], "Lead Source": r["lead_source"],
                         "Technology": r["technology"], "Lead Type": r["lead_type"],
                         "Source Sheet": r["tab"], "Sheet Row": m["row"], "Flags": " | ".join(f)})

    # month by month (historical validation)
    tabs_l = list(per_tab)
    months = sorted({_month_of(r["date"]) for r in kept if r["date"]} | {_month_of(r["date"]) for r in kept_old if r["date"]})
    mbm = []
    for ms in months:
        row = {"Month": ms.strftime("%b-%Y")}
        tn = to = 0
        for t in tabs_l:
            n = sum(1 for r in kept if r["tab"] == t and _month_of(r["date"]) == ms)
            o = sum(1 for r in kept_old if r["tab"] == t and _month_of(r["date"]) == ms)
            row[f"{t} (corrected)"], row[f"{t} (previous)"] = n, o
            tn += n; to += o
        row["Total (corrected)"], row["Total (previous)"], row["Change"] = tn, to, tn - to
        mbm.append(row)

    text_dates = [{"Source Sheet": r["tab"], "Sheet Row": info[id(r)]["row"], "Raw Timestamp": info[id(r)]["raw_date"],
                   "Date rule": info[id(r)]["rule"], "Read as": r["date"], "Previous logic": info[id(r)]["old"]}
                  for r in new_all if info[id(r)]["kind"] in ("text", "unparsed")]
    other_diff = [{"Source Sheet": r["tab"], "Sheet Row": info[id(r)]["row"], "Lead Name": r["name"],
                   "Timestamp date": r["date"], "Column": h, "Column date": od}
                  for r in counted for h, od in info[id(r)]["other"].items()
                  if od and _month_of(od) != _month_of(r["date"])]
    return {"month": f"{month_start:%b-%Y}", "counted": len(counted), "rows": rows_out, "readings": readings,
            "per_tab": {t: {k: v for k, v in pt.items() if k != "recs"} for t, pt in per_tab.items()},
            "month_by_month": mbm, "text_dates": text_dates, "other_diff": other_diff,
            "by_tab": collections.Counter(r["tab"] for r in counted),
            "by_source": collections.Counter(r["lead_source"] for r in counted),
            "total_records": len(new_all), "after_dedupe": len(kept), "removed_total": removed}


def write_xlsx(path, res, meta):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    NAVY = PatternFill("solid", fgColor="1B355E")
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Summary"
    L = [("Walk-In trace", res["month"]),
         ("Counted by the report (DATE RULE)", res["counted"]),
         ("Spreadsheet locale / time zone", f"{meta.get('locale', '?')} / {meta.get('timeZone', '?')}"), ("", "")]
    for tab, pt in res["per_tab"].items():
        L += [(f"{tab}", ""),
              ("   Timestamp cells", ", ".join(f"{k} {v}" for k, v in pt["types"].most_common()))]
        if pt["formats"]:
            L.append(("   Date formats in the column", ", ".join(f"'{k}' {v}" for k, v in pt["formats"].most_common(5))))
        L.append(("   Most common display shapes", ", ".join(f"{k} ×{v}" for k, v in pt["shapes"])))
        L.append((f"   {res['month']} count by reading", ""))
        L += [(f"      {k}", v) for k, v in res["readings"][tab].items()]
        L.append(("", ""))
    L += [("By lead source (counted)", "")] + [(f"   {k}", v) for k, v in res["by_source"].most_common()]
    L += [("", ""), ("All records (both tabs)", res["total_records"]), ("After de-duplication", res["after_dedupe"])]
    for a, b in L:
        ws.append([a, b])
    ws.column_dimensions["A"].width = 58; ws.column_dimensions["B"].width = 80
    for name, data in ((f"Counted {res['month']} ({res['counted']})", res["rows"]),
                       ("Month by month", res["month_by_month"]),
                       ("Text dates", res["text_dates"]),
                       ("Other date columns", res["other_diff"])):
        s = wb.create_sheet(name[:31])
        if not data:
            s.append(["(none)"]); continue
        cols = list(data[0])
        s.append(cols)
        for c in s[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = NAVY
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for d in data:
            s.append([d[c] for c in cols])
        for i, c in enumerate(cols, 1):
            s.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(60, max(11, len(c) + 3))
        for row in s.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, dt.date):
                    c.number_format = "DD-MMM-YYYY"
                if cols[c.column - 1] in ("Flags", "Change") and c.value:
                    c.fill = PatternFill("solid", fgColor="FCF3CF")
        s.freeze_panes = "A2"
    wb.save(path)


def _column_formats(sheets, sid, tab, col_idx):
    """Counter of the number-format patterns used by the Timestamp column."""
    try:
        col = openpyxl_col(col_idx + 1)
        resp = sheets.spreadsheets().get(
            spreadsheetId=sid, ranges=[f"'{tab}'!{col}2:{col}"], includeGridData=True,
            fields="sheets.data.rowData.values(effectiveValue,effectiveFormat.numberFormat)").execute()
        out = collections.Counter()
        for rd in resp["sheets"][0]["data"][0].get("rowData", []):
            for v in rd.get("values", []):
                ev = v.get("effectiveValue", {})
                if not ev:
                    continue
                nf = v.get("effectiveFormat", {}).get("numberFormat", {})
                out[(nf.get("pattern") or nf.get("type") or "(none)") + ("" if "numberValue" in ev else " [text]")] += 1
        return out
    except Exception as e:                                  # evidence only — never fatal
        return collections.Counter({f"(could not read formats: {type(e).__name__})": 1})


def openpyxl_col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26); s = chr(65 + r) + s
    return s


def main(argv=None):
    ap = argparse.ArgumentParser(description="Trace one month's Walk-In count to its source rows (read-only).")
    ap.add_argument("--month", default="2026-05", help="YYYY-MM (default 2026-05)")
    a = ap.parse_args(argv)
    ms = dt.datetime.strptime(a.month, "%Y-%m").date()
    import google_utils
    sheets, _drive = google_utils.get_services()
    sid = cfg.get("source.spreadsheet_id", "")
    tab_names = cfg.get("source.tabs", ["Walk-In New", "Walk-In Old"])
    meta = sheets.spreadsheets().get(spreadsheetId=sid, fields="properties(locale,timeZone)").execute().get("properties", {})
    tabs = {}
    for t in tab_names:
        val = google_utils.read_tab(sheets, sid, t, date_render="SERIAL_NUMBER")
        disp = google_utils.read_tab(sheets, sid, t, date_render="FORMATTED_STRING")
        which = "new" if "new" in t.lower() else "old"
        di = WD._col_index(val[0] if val else [], WD._mapping()[which]["date"])
        tabs[t] = {"val": val, "disp": disp,
                   "formats": _column_formats(sheets, sid, t, di) if di is not None else None}
    res = trace(tabs, ms)
    print(f"\nSpreadsheet locale / time zone: {meta.get('locale')} / {meta.get('timeZone')}")
    print(f"{res['month']}: {res['counted']} Walk-In leads counted by the report (DATE RULE)")
    for tab, pt in res["per_tab"].items():
        print(f"\n  {tab}: Timestamp cells — " + ", ".join(f"{k} {v}" for k, v in pt["types"].most_common()))
        if pt["formats"]:
            print("     formats: " + ", ".join(f"'{k}' {v}" for k, v in pt["formats"].most_common(4)))
        for k, v in res["readings"][tab].items():
            print(f"     {k:<48} {v}")
    changed = [m for m in res["month_by_month"] if m["Change"]]
    print(f"\nMonths whose count changed vs the previous logic: {len(changed)}"
          + "".join(f"\n   {m['Month']}: {m['Total (previous)']} -> {m['Total (corrected)']}" for m in changed))
    paths.ensure_dirs()
    out = os.path.join(str(paths.OUTPUT_DIR), f"_trace - Walk-Ins {res['month']}.xlsx")
    write_xlsx(out, res, meta)
    print(f"\nSaved: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
